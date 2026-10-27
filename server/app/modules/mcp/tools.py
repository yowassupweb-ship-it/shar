"""Инструменты MCP «Офиса»: что агент может делать с заметками, документами, таблицами, презентациями и папками.

Каждый инструмент зовёт те же обработчики, что и интерфейс (workspace/router.py, workspace/folders.py), от имени
владельца ключа: права, история версий, живые события коллегам — те же. Сверх того сообщаем и самому человеку:
интерфейс не шлёт автору его же правку, а правка агента для человека — чужая (открытая заметка должна её увидеть).

Форматы для агента: заметка — Markdown (вкладки — строками `<!-- shar-tab: Название -->`), документ — Markdown
(перевод в снимок редактора и обратно — docformat.py), таблица — строки ячеек, презентация — только текст.
"""

from __future__ import annotations

import base64
import io
import json
import logging
from collections.abc import Callable
from typing import Any

from fastapi import HTTPException

from app.core.security import Identity
from app.modules.mcp import docformat
from app.modules.workspace import folders as ws_folders
from app.modules.workspace import router as ws
from app.modules.workspace import store

log = logging.getLogger("shar.mcp")

KIND_NAMES = {"note": "заметка", "document": "документ", "table": "таблица", "presentation": "презентация"}
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 50


class ToolError(Exception):
    """Ошибка, которую агент должен увидеть текстом (isError), а не сбоем протокола."""


def _call(fn: Callable[..., Any], *args: Any) -> Any:
    try:
        return fn(*args)
    except HTTPException as exc:
        raise ToolError(str(exc.detail)) from exc
    except store.WorkspaceError as exc:
        raise ToolError(exc.message) from exc


def _notify_self(me: Identity, item_id: str, kind: str, *, content: bool) -> None:
    """Правка агента — и для самого человека событие: список перестроится, открытый материал подтянет правку."""
    try:
        from app.core.realtime import sse_manager

        live = store.get_row(item_id, full=False)
        if not live:
            return
        if content:
            sse_manager.push(me.user_id, {"type": "workspace_changed", "reason": "content", "itemId": item_id, "kind": kind,
                                          "updatedAt": live["updated_at"], "by": "Агент MCP", "byId": "mcp"})
        store.announce({me.user_id}, live, "changed")
    except Exception as exc:  # noqa: BLE001 — событие лишь удобство, не повод ломать ответ агенту
        log.warning("mcp: could not notify %s about %s: %s", me.user_id, item_id, exc)


def _folder_paths(me: Identity) -> tuple[list[dict[str, Any]], dict[str, str]]:
    folders = ws_folders._folders(me.user_id)
    by_id = {f["id"]: f for f in folders}

    def path(f: dict[str, Any], depth: int = 0) -> str:
        parent = by_id.get(f.get("parentId") or "")
        return f"{path(parent, depth + 1)} / {f['title']}" if parent and depth < 16 else f["title"]

    return folders, {f["id"]: path(f) for f in folders}


def _item_brief(item: dict[str, Any], paths: dict[str, str]) -> dict[str, Any]:
    return {
        "id": item["id"], "kind": item["kind"], "title": item["title"] or "Без названия",
        "folderId": item.get("folderId"), "folder": paths.get(item.get("folderId") or "", ""),
        "updatedAt": item["updatedAt"], "owner": item.get("ownerName") or "", "canEdit": bool(item.get("canEdit")),
        "snippet": item.get("snippet") or "",
    }


def _place(me: Identity, item_id: str, folder_id: str | None) -> None:
    if folder_id:
        _call(ws_folders.place_item, item_id, ws_folders.PlaceBody(folderId=folder_id), me)


def _table_to_rows(content: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(base64.b64decode(content)), read_only=True, data_only=True)
    sheets = []
    for sheet in book.worksheets:
        rows: list[list[Any]] = []
        for row in sheet.iter_rows(max_row=MAX_TABLE_ROWS, max_col=MAX_TABLE_COLS, values_only=True):
            rows.append(["" if v is None else (v if isinstance(v, (int, float, bool)) else str(v)) for v in row])
        while rows and not any(str(v).strip() for v in rows[-1]):
            rows.pop()
        # Пустые столбцы справа не нужны агенту (read_only отдаёт строки до max_col).
        width = max((max((i + 1 for i, v in enumerate(r) if str(v).strip()), default=0) for r in rows), default=0)
        rows = [r[:width] for r in rows]
        sheets.append({"name": sheet.title, "rows": rows, "truncated": (sheet.max_row or 0) > MAX_TABLE_ROWS})
    return sheets


def _rows_to_table(sheets: list[dict[str, Any]]) -> str:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for index, sheet in enumerate(sheets or [{"name": "Лист1", "rows": []}]):
        ws_sheet = book.create_sheet(title=(str(sheet.get("name") or f"Лист{index + 1}"))[:31])
        for row in (sheet.get("rows") or [])[:MAX_TABLE_ROWS * 4]:
            ws_sheet.append([cell if isinstance(cell, (int, float, bool)) or cell is None else str(cell) for cell in list(row)[:MAX_TABLE_COLS * 4]])
    out = io.BytesIO()
    book.save(out)
    return base64.b64encode(out.getvalue()).decode("ascii")


# ── инструменты ─────────────────────────────────────────────────────────────

def list_items(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    kind = args.get("kind") or None
    if kind and kind not in store.KINDS:
        raise ToolError("kind: note, document, table или presentation")
    items = _call(ws.list_items, kind, str(args.get("query") or ""), False, 1000, me)["items"]
    _, paths = _folder_paths(me)
    folder_id = args.get("folder_id")
    if folder_id:
        items = [i for i in items if (i.get("folderId") or "") == str(folder_id)]
    limit = max(1, min(int(args.get("limit") or 50), 200))
    return {"items": [_item_brief(i, paths) for i in items[:limit]], "total": len(items)}


def list_folders(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    folders, paths = _folder_paths(me)
    return {"folders": [{"id": f["id"], "title": f["title"], "parentId": f.get("parentId"), "path": paths[f["id"]],
                         "counts": f.get("counts") or {}, "mine": f.get("role") == "owner"} for f in folders]}


def read_item(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    item = _call(ws.get_item, str(args.get("id") or ""), me)["item"]
    content = item.get("content") or ""
    out: dict[str, Any] = {"id": item["id"], "kind": item["kind"], "title": item["title"], "updatedAt": item["updatedAt"],
                           "canEdit": bool(item.get("canEdit"))}
    if item["kind"] == "note":
        out.update(format="markdown", content=content)
    elif item["kind"] == "document":
        try:
            out.update(format="markdown", content=docformat.snapshot_to_markdown(json.loads(content)))
        except ValueError:
            out.update(format="markdown", content="")
    elif item["kind"] == "table":
        try:
            out.update(format="sheets", sheets=_table_to_rows(content))
        except Exception as exc:
            raise ToolError(f"Таблицу не удалось прочитать: {exc}") from exc
    else:
        out.update(format="text", content=store.text_of(item["kind"], content))
    return out


def create_note(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    item = _call(ws.create_item, ws.ItemCreate(kind="note", title=str(args.get("title") or ""), content=str(args.get("markdown") or "")), me)["item"]
    _place(me, item["id"], args.get("folder_id"))
    _notify_self(me, item["id"], "note", content=False)
    return {"id": item["id"], "kind": "note", "title": item["title"]}


def update_note(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    item_id = str(args.get("id") or "")
    current = _call(ws.get_item, item_id, me)["item"]
    if current["kind"] != "note":
        raise ToolError(f"Это {KIND_NAMES[current['kind']]}, а не заметка")
    text = str(args.get("markdown") or "")
    if (args.get("mode") or "replace") == "append":
        base = current.get("content") or ""
        text = f"{base.rstrip()}\n\n{text}" if base.strip() else text
    patch = ws.ItemPatch(content=text, **({"title": str(args["title"])} if args.get("title") else {}))
    _call(ws.patch_item, item_id, patch, me)
    _notify_self(me, item_id, "note", content=True)
    return {"id": item_id, "ok": True}


def create_document(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    title = str(args.get("title") or "")
    snapshot = docformat.markdown_to_snapshot(title, str(args.get("markdown") or ""))
    item = _call(ws.create_item, ws.ItemCreate(kind="document", title=title, content=json.dumps(snapshot, ensure_ascii=False)), me)["item"]
    _place(me, item["id"], args.get("folder_id"))
    _notify_self(me, item["id"], "document", content=False)
    return {"id": item["id"], "kind": "document", "title": item["title"]}


def update_document(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    item_id = str(args.get("id") or "")
    current = _call(ws.get_item, item_id, me)["item"]
    if current["kind"] != "document":
        raise ToolError(f"Это {KIND_NAMES[current['kind']]}, а не документ")
    markdown = str(args.get("markdown") or "")
    if (args.get("mode") or "append") == "append":
        try:
            snapshot = json.loads(current.get("content") or "")
        except ValueError:
            snapshot = docformat.empty_snapshot(current["title"])
        snapshot = docformat.append_markdown(snapshot, markdown)
    else:
        snapshot = docformat.markdown_to_snapshot(current["title"], markdown)
    _call(ws.patch_item, item_id, ws.ItemPatch(content=json.dumps(snapshot, ensure_ascii=False)), me)
    _notify_self(me, item_id, "document", content=True)
    return {"id": item_id, "ok": True}


def create_table(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    sheets = args.get("sheets") or [{"name": "Лист1", "rows": args.get("rows") or []}]
    if not isinstance(sheets, list):
        raise ToolError("sheets — список листов {name, rows}")
    content = _rows_to_table(sheets)
    item = _call(ws.create_item, ws.ItemCreate(kind="table", title=str(args.get("title") or ""), content=content), me)["item"]
    _place(me, item["id"], args.get("folder_id"))
    _notify_self(me, item["id"], "table", content=False)
    return {"id": item["id"], "kind": "table", "title": item["title"]}


def rename_item(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    item_id = str(args.get("id") or "")
    title = str(args.get("title") or "").strip()
    if not title:
        raise ToolError("Нужно название")
    item = _call(ws.patch_item, item_id, ws.ItemPatch(title=title), me)["item"]
    _notify_self(me, item_id, item["kind"], content=True)
    return {"id": item_id, "title": item["title"]}


def move_item(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    item_id = str(args.get("id") or "")
    folder_id = args.get("folder_id") or None
    _call(ws_folders.place_item, item_id, ws_folders.PlaceBody(folderId=str(folder_id) if folder_id else None), me)
    return {"id": item_id, "folderId": folder_id}


def create_folder(me: Identity, args: dict[str, Any]) -> dict[str, Any]:
    parent = args.get("parent_id") or None
    body = ws_folders.FolderBody(title=str(args.get("title") or "").strip() or "Новая папка", parentId=str(parent) if parent else None)
    folder = _call(ws_folders.create_folder, body, me)["folder"]
    return {"id": folder["id"], "title": folder["title"], "parentId": folder.get("parentId")}


_ID = {"type": "string", "description": "id материала (из list_items)"}
_FOLDER = {"type": "string", "description": "id папки (из list_folders); не указан — вне папок"}

TOOLS: list[dict[str, Any]] = [
    {"name": "list_items", "handler": list_items, "readOnly": True,
     "description": "Материалы «Офиса», доступные человеку: заметки, документы, таблицы, презентации. Поиск по названию и тексту.",
     "inputSchema": {"type": "object", "properties": {
         "kind": {"type": "string", "enum": list(store.KINDS)},
         "query": {"type": "string", "description": "слова из названия или текста"},
         "folder_id": _FOLDER, "limit": {"type": "integer", "minimum": 1, "maximum": 200, "default": 50}}}},
    {"name": "list_folders", "handler": list_folders, "readOnly": True,
     "description": "Папки человека (свои и открытые ему) с путём и числом материалов каждого вида.",
     "inputSchema": {"type": "object", "properties": {}}},
    {"name": "read_item", "handler": read_item, "readOnly": True,
     "description": "Содержимое материала. Заметка и документ — Markdown, таблица — листы со строками ячеек (до 500 строк), презентация — текст слайдов.",
     "inputSchema": {"type": "object", "properties": {"id": _ID}, "required": ["id"]}},
    {"name": "create_note", "handler": create_note,
     "description": "Создать заметку из Markdown. Вкладки — строкой `<!-- shar-tab: Название -->` перед текстом вкладки.",
     "inputSchema": {"type": "object", "properties": {"title": {"type": "string"}, "markdown": {"type": "string"}, "folder_id": _FOLDER},
                     "required": ["title", "markdown"]}},
    {"name": "update_note", "handler": update_note,
     "description": "Изменить заметку: mode=replace заменяет текст целиком (служебные строки вкладок сохраняйте), mode=append дописывает в конец. Прежний текст остаётся в истории версий.",
     "inputSchema": {"type": "object", "properties": {"id": _ID, "markdown": {"type": "string"}, "mode": {"type": "string", "enum": ["replace", "append"], "default": "replace"},
                                                     "title": {"type": "string", "description": "новое название (необязательно)"}}, "required": ["id", "markdown"]}},
    {"name": "create_document", "handler": create_document,
     "description": "Создать документ (листы A4) из Markdown: заголовки #–###, **жирный**, *курсив*, `код`, ссылки, списки.",
     "inputSchema": {"type": "object", "properties": {"title": {"type": "string"}, "markdown": {"type": "string"}, "folder_id": _FOLDER},
                     "required": ["title", "markdown"]}},
    {"name": "update_document", "handler": update_document,
     "description": "Изменить документ: mode=append дописывает Markdown в конец, сохраняя оформление (по умолчанию); mode=replace пересобирает документ из Markdown — сложное оформление, таблицы и картинки пропадут. Прежняя версия остаётся в истории версий.",
     "inputSchema": {"type": "object", "properties": {"id": _ID, "markdown": {"type": "string"}, "mode": {"type": "string", "enum": ["append", "replace"], "default": "append"}},
                     "required": ["id", "markdown"]}},
    {"name": "create_table", "handler": create_table,
     "description": "Создать таблицу. rows — строки ячеек одного листа, или sheets — несколько листов {name, rows}. Первая строка обычно — заголовки.",
     "inputSchema": {"type": "object", "properties": {
         "title": {"type": "string"}, "folder_id": _FOLDER,
         "rows": {"type": "array", "items": {"type": "array", "items": {}}},
         "sheets": {"type": "array", "items": {"type": "object", "properties": {"name": {"type": "string"}, "rows": {"type": "array", "items": {"type": "array", "items": {}}}}}}},
         "required": ["title"]}},
    {"name": "rename_item", "handler": rename_item, "description": "Переименовать материал.",
     "inputSchema": {"type": "object", "properties": {"id": _ID, "title": {"type": "string"}}, "required": ["id", "title"]}},
    {"name": "move_item", "handler": move_item, "description": "Положить материал в свою папку или вынуть из папки (folder_id не указан).",
     "inputSchema": {"type": "object", "properties": {"id": _ID, "folder_id": _FOLDER}, "required": ["id"]}},
    {"name": "create_folder", "handler": create_folder, "description": "Создать свою папку (parent_id — внутри другой своей папки).",
     "inputSchema": {"type": "object", "properties": {"title": {"type": "string"}, "parent_id": {"type": "string"}}, "required": ["title"]}},
]

BY_NAME = {t["name"]: t for t in TOOLS}


def describe() -> list[dict[str, Any]]:
    """Список инструментов для tools/list (без обработчиков)."""
    return [{"name": t["name"], "description": t["description"], "inputSchema": t["inputSchema"],
             "annotations": {"readOnlyHint": bool(t.get("readOnly")), "destructiveHint": False}} for t in TOOLS]
