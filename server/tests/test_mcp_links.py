"""Ссылки ячеек в MCP update_table: смена значения сбрасывает старую привязку, ссылку можно задать явно."""

import base64
import io

from openpyxl import Workbook, load_workbook

from app.modules.mcp import tools


def _book():
    book = Workbook()
    sheet = book.active
    sheet["F11"] = "Старый документ"
    sheet["F11"].hyperlink = "https://docs.google.com/document/d/1_q/edit"
    return book, sheet


def _roundtrip(book):
    out = io.BytesIO()
    book.save(out)
    return load_workbook(io.BytesIO(out.getvalue()))


def test_new_value_drops_the_old_link():
    book, sheet = _book()
    tools._write_cell(sheet, "F11", {"ref": "F11", "value": None})
    tools._write_cell(sheet, "F11", {"ref": "F11", "value": "https://vokrug-sveta.shar-os.ru/office?kind=document&id=94"})
    cell = _roundtrip(book).active["F11"]
    assert cell.value == "https://vokrug-sveta.shar-os.ru/office?kind=document&id=94"
    assert cell.hyperlink is None  # привязка к docs.google.com не пережила смену текста


def test_explicit_link_in_three_forms():
    book, sheet = _book()
    tools._write_cell(sheet, "A1", {"ref": "A1", "value": "Документ", "link": "/office?kind=document&id=94"})
    tools._write_cell(sheet, "A2", {"ref": "A2", "value": {"text": "Таблица", "hyperlink": "/office?kind=table&id=103"}})
    tools._write_cell(sheet, "A3", {"ref": "A3", "value": '=HYPERLINK("/office?kind=note&id=7"; "Заметка")'})
    tools._write_cell(sheet, "A4", {"ref": "A4", "link": "https://example.com"})  # без текста — текстом станет адрес
    reloaded = _roundtrip(book).active
    assert (reloaded["A1"].value, reloaded["A1"].hyperlink.target) == ("Документ", "/office?kind=document&id=94")
    assert (reloaded["A2"].value, reloaded["A2"].hyperlink.target) == ("Таблица", "/office?kind=table&id=103")
    assert (reloaded["A3"].value, reloaded["A3"].hyperlink.target) == ("Заметка", "/office?kind=note&id=7")
    assert (reloaded["A4"].value, reloaded["A4"].hyperlink.target) == ("https://example.com", "https://example.com")


def test_read_gives_links_and_item_url():
    book, _sheet = _book()
    out = io.BytesIO()
    book.save(out)
    sheets = tools._table_rows(base64.b64encode(out.getvalue()).decode())
    assert sheets[0]["links"] == {"F11": "https://docs.google.com/document/d/1_q/edit"}
    assert tools.item_url("document", 94) == "/office?kind=document&id=94"
